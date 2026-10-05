"""Open real QA exports with independent document readers (QA packages only)."""
import io
import json
import os
import zipfile
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

root = Path(os.environ["MALIK_QA_OUTPUT"])
doc = Document(root / "export.docx")
assert "Алматы" in "\n".join(p.text for p in doc.paragraphs) + str([[c.text for c in row.cells] for table in doc.tables for row in table.rows])
print("PASS DOCX: opened with python-docx; Cyrillic preserved")
book = load_workbook(root / "export.xlsx")
assert book.active["A2"].value == "Алматы" and book.active["B2"].value == 12
print("PASS XLSX: opened with openpyxl; text and numeric cells preserved")
pdf = PdfReader(root / "export.pdf")
# PDF text readers may insert line breaks between separately positioned words.
assert "Кириллица: Алматы, Астана" in " ".join(pdf.pages[0].extract_text().split())
assert pdf.pages[0]["/Annots"]
print("PASS PDF: pypdf extracts Cyrillic and link annotations")
table = PdfReader(root / "table.pdf")
assert len(table.pages) > 1 and all("Заголовок" in page.extract_text() for page in table.pages)
print(f"PASS paginated PDF: {len(table.pages)} pages, repeated table headers")
with zipfile.ZipFile(root / "export.pptx") as deck:
    slides = [name for name in deck.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml")]
    assert len(slides) == 2 and "План Алматы" in deck.read(slides[0]).decode("utf-8")
    assert "без изображений" in deck.read(slides[0]).decode("utf-8")
print("PASS PPTX: real OOXML package; 2 slides and honest image-omission caption")
with zipfile.ZipFile(root / "export.zip") as package:
    assert len(package.namelist()) == 8
    assert all(name.endswith((".docx", ".pdf", ".xlsx", ".csv", ".md", ".txt", ".html", ".json")) for name in package.namelist())
    for name in package.namelist():
        if name.endswith(".docx"): Document(io.BytesIO(package.read(name)))
        if name.endswith(".xlsx"): load_workbook(io.BytesIO(package.read(name)))
        if name.endswith(".json"): json.loads(package.read(name))
print("PASS ZIP: all 8 actual formats, nested documents reopen")
print("6/6 independent file checks passed")
